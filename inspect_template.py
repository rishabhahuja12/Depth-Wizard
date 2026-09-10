import pptx

prs = pptx.Presentation(r"E:\rishabh\sih\SIH2026-IDEA-Presentation-Format (1).pptx")
print(f"Slide count: {len(prs.slides)}")
print(f"Dimensions: {prs.slide_width.inches:.2f} x {prs.slide_height.inches:.2f} inches")

for idx, s in enumerate(prs.slides):
    print(f"\n=== SLIDE {idx+1} ===")
    for shp in s.shapes:
        txt = shp.text.strip().replace('\n', ' ') if shp.has_text_frame else ''
        print(f"  Shape '{shp.name}' (type={shp.shape_type}): left={shp.left.inches:.2f}, top={shp.top.inches:.2f}, w={shp.width.inches:.2f}, h={shp.height.inches:.2f} | text='{txt[:70]}'")
