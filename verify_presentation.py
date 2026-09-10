import pptx

prs = pptx.Presentation(r"E:\rishabh\sih\DepthWizard_SIH2026_Final_Presentation.pptx")
print(f"Total Slides: {len(prs.slides)}")
assert len(prs.slides) == 6, f"Expected 6 slides, found {len(prs.slides)}"

for idx, slide in enumerate(prs.slides):
    pictures = [s for s in slide.shapes if s.shape_type == pptx.enum.shapes.MSO_SHAPE_TYPE.PICTURE]
    text_shapes = [s for s in slide.shapes if s.has_text_frame and s.text.strip()]
    title = ""
    for s in text_shapes:
        if "PROPOSED" in s.text or "TECHNICAL" in s.text or "FEASIBILITY" in s.text or "DISASTER" in s.text or "RESEARCH" in s.text or "DEPTHWIZARD" in s.text:
            title = s.text.split("\n")[0][:60]
            break
    print(f"Slide {idx + 1}: Title='{title}' | Pictures={len(pictures)} | Text/Cards={len(text_shapes)}")
    for p in pictures:
        print(f"   -> Image: left={p.left.inches:.2f}\", top={p.top.inches:.2f}\", w={p.width.inches:.2f}\", h={p.height.inches:.2f}\"")

print("\nAll 6 slides verified successfully!")
